"""Build two generic architecture slides with original icons and editable connections."""

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
ICONS = ASSETS / "architecture-icons"
NAVY = "16304E"
TEAL = "007F83"
BLUE = "2866BD"
MUTED = "52677D"
BORDER = "CFDAE5"
PALE = "F2F6FA"
WHITE = "FFFFFF"
FONT = "Segoe UI"

RUN = "https://learn.microsoft.com/en-us/rest/api/fabric/core/job-scheduler/run-on-demand-item-job"
STATUS = "https://learn.microsoft.com/en-us/rest/api/fabric/core/job-scheduler/get-item-job-instance"
WAREHOUSE = "https://docs.getdbt.com/docs/local/connect-data-platform/fabric-setup"
SPARK = "https://docs.getdbt.com/docs/local/connect-data-platform/fabricspark-setup"
WEB = "https://learn.microsoft.com/en-us/fabric/data-factory/web-activity"


def rgb(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill=WHITE, line=BORDER, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(1)
    if rounded:
        shape.adjustments[0] = 0.12
    return shape


def text(slide, content, x, y, w, h, size=16, color=NAVY, bold=False, align=None, url=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_top = frame.margin_bottom = 0
    frame.margin_left = frame.margin_right = 0
    for i, line in enumerate(content.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        if align is not None:
            p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        if url:
            run.hyperlink.address = url
    return shape


def icon(name):
    path = ICONS / f"{name}.png"
    image = Image.new("RGBA", (256, 256), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    color = "#" + (TEAL if name in ("cloud", "database", "copy", "gate") else BLUE)
    draw.rounded_rectangle((8, 8, 248, 248), radius=48, fill=color)
    ink = "#FFFFFF"
    if name == "cloud":
        draw.ellipse((54, 100, 132, 178), fill=ink)
        draw.ellipse((86, 61, 181, 164), fill=ink)
        draw.ellipse((144, 104, 215, 178), fill=ink)
        draw.rectangle((86, 124, 181, 178), fill=ink)
    elif name == "database":
        draw.rectangle((59, 79, 197, 174), fill=ink)
        draw.ellipse((59, 149, 197, 199), fill=ink)
        draw.ellipse((59, 54, 197, 104), fill=ink, outline=color, width=6)
        draw.arc((59, 94, 197, 144), 0, 180, fill=color, width=6)
    elif name == "airflow":
        points = [(65, 73), (191, 73), (65, 183), (191, 183)]
        for p in points:
            draw.line((128, 128, *p), fill=ink, width=9)
            draw.ellipse((p[0] - 15, p[1] - 15, p[0] + 15, p[1] + 15), fill=ink)
        draw.ellipse((103, 103, 153, 153), fill=ink)
    elif name == "runner":
        draw.rounded_rectangle((43, 61, 213, 195), radius=12, outline=ink, width=9)
        draw.line((68, 97, 94, 122, 68, 148), fill=ink, width=9)
        draw.line((116, 150, 177, 150), fill=ink, width=9)
    elif name == "api":
        draw.line((93, 76, 52, 128, 93, 180), fill=ink, width=11)
        draw.line((163, 76, 204, 128, 163, 180), fill=ink, width=11)
        draw.line((143, 69, 113, 187), fill=ink, width=9)
    elif name == "copy":
        draw.rounded_rectangle((48, 53, 149, 167), radius=9, outline=ink, width=9)
        draw.rounded_rectangle((106, 105, 207, 219), radius=9, fill=color, outline=ink, width=9)
    elif name == "clock":
        draw.ellipse((53, 53, 203, 203), outline=ink, width=10)
        draw.line((128, 78, 128, 128, 166, 146), fill=ink, width=10)
    elif name == "gate":
        draw.line((59, 128, 105, 174, 196, 80), fill=ink, width=17)
    else:
        raise ValueError(f"Unknown icon: {name}")
    image.resize((192, 192), Image.Resampling.LANCZOS).save(path)
    return path


def picture(slide, name, x, y, size=0.58):
    return slide.shapes.add_picture(str(ICONS / f"{name}.png"), Inches(x), Inches(y), Inches(size), Inches(size))


def card(slide, x, y, w, h, heading, detail, image, size=18):
    box(slide, x, y, w, h)
    picture(slide, image, x + 0.15, y + (h - 0.58) / 2)
    text(slide, heading, x + 0.88, y + 0.15, w - 1.02, 0.36, size=size, bold=True)
    text(slide, detail, x + 0.88, y + 0.57, w - 1.02, max(0.3, h - 0.64), size=12, color=MUTED)


def route(slide, points, color=TEAL, dashed=False, both=False):
    for i, (start, end) in enumerate(zip(points, points[1:])):
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(start[0]), Inches(start[1]), Inches(end[0]), Inches(end[1]),
        )
        line.line.color.rgb = rgb(color)
        line.line.width = Pt(2.1)
        ln = line._element.spPr.get_or_add_ln()
        if dashed:
            dash = OxmlElement("a:prstDash")
            dash.set("val", "dash")
            ln.append(dash)
        if i == len(points) - 2:
            arrow = OxmlElement("a:tailEnd")
            arrow.set("type", "triangle")
            ln.append(arrow)
        if both and i == 0:
            arrow = OxmlElement("a:headEnd")
            arrow.set("type", "triangle")
            ln.append(arrow)


def label(slide, content, x, y, w, color=TEAL, size=12):
    shape = box(slide, x, y, w, 0.32, WHITE, WHITE, rounded=False)
    shape.line.fill.background()
    text(slide, content, x, y + 0.025, w, 0.28, size=size, color=color, bold=True, align=PP_ALIGN.CENTER)


def panel(slide, x, w, heading, image):
    box(slide, x, 2.03, w, 4.98, fill=PALE)
    picture(slide, image, x + 0.22, 2.22, 0.36)
    text(slide, heading, x + 0.73, 2.26, w - 0.95, 0.3, size=12, bold=True, color=MUTED)


def base(prs, number, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    box(slide, 0, 0, 16, 0.12, fill=TEAL, line=TEAL, rounded=False)
    text(slide, "FABRIC  /  EXTERNAL ORCHESTRATION", 0.6, 0.38, 14, 0.28, size=11, color=TEAL, bold=True)
    text(slide, title, 0.6, 0.85, 14.8, 0.62, size=31, bold=True)
    text(slide, subtitle, 0.62, 1.52, 14.6, 0.35, size=16, color=MUTED)
    text(slide, f"{number} / 2", 14.65, 8.53, 0.75, 0.22, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def footer(slide, title, detail, links):
    box(slide, 0.6, 7.25, 14.8, 0.93, fill="EAF4F4", line="D7E9E9")
    text(slide, title, 0.85, 7.39, 14.3, 0.32, size=18, bold=True)
    text(slide, detail, 0.85, 7.83, 14.3, 0.25, size=12, color=MUTED)
    x = 0.62
    for caption, url, width in links:
        text(slide, caption, x, 8.53, width, 0.24, size=10, color=BLUE, url=url)
        x += width + 0.25


def slide_external(prs):
    slide = base(
        prs, 1, "External Airflow orchestrates Fabric",
        "Customer-managed scheduling; Fabric performs ingestion and transformations.",
    )
    panel(slide, 0.5, 5.0, "CUSTOMER-MANAGED HOST", "runner")
    panel(slide, 8.5, 7.0, "MICROSOFT FABRIC", "cloud")
    route(slide, [(3.0, 3.87), (3.0, 5.25)])
    route(slide, [(12.0, 3.87), (12.0, 5.25)])
    route(slide, [(5.15, 3.08), (9.0, 3.08)])
    route(slide, [(5.15, 3.62), (9.0, 3.62)], color=BLUE, dashed=True, both=True)
    route(slide, [(5.15, 5.78), (9.0, 5.78)])
    card(slide, 0.85, 2.79, 4.3, 1.08, "Airflow DAG", "Schedule / dependencies / retries", "airflow")
    card(slide, 0.85, 5.25, 4.3, 1.08, "External dbt runner", "Approved runtime + pinned project", "runner")
    card(slide, 9.0, 2.79, 6.0, 1.08, "Fabric pipeline + job API", "Ingestion and tracked job execution", "cloud")
    card(slide, 9.0, 5.25, 6.0, 1.08, "Warehouse OR Lakehouse / Spark", "Bronze -> dbt-managed Silver / Gold", "database", size=17)
    label(slide, "1  POST: start ingestion", 5.6, 2.65, 2.95)
    label(slide, "2  GET: poll job status", 5.6, 3.88, 2.95, color=BLUE)
    label(slide, "3  Run dbt after ingestion completes", 0.95, 4.36, 4.1, size=11.5)
    label(slide, "Copy into Bronze", 10.3, 4.36, 3.4)
    label(slide, "4  SQL or Spark/Livy", 5.6, 5.32, 2.95)
    text(slide, "On premises or a cloud host", 0.9, 6.62, 4.2, 0.23, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    text(slide, "Solid: submit / execute     Dashed: status request / response", 8.95, 6.62, 6.1, 0.23, size=10.5, color=MUTED)
    footer(
        slide,
        "Airflow is the workflow owner; Fabric supplies the processing.",
        "Warehouse: dbt-fabric over SQL  |  Lakehouse: dbt-fabricspark over Spark/Livy  |  Configure Entra identity + network access.",
        [("Start a Fabric job", RUN, 1.65), ("Monitor its status", STATUS, 1.65),
         ("Warehouse adapter", WAREHOUSE, 1.8), ("Lakehouse adapter", SPARK, 1.8)],
    )
    slide.notes_slide.notes_text_frame.text = (
        "Talk track: The customer-managed host can be on premises or in another cloud; it is not "
        "necessarily a laptop or a single server. Airflow owns the overall DAG. It authenticates to "
        "Fabric, submits an ingestion pipeline, saves the job instance Location, and polls until "
        "Completed. The dashed arrow represents client-initiated status checks and responses, not "
        "Fabric pushing a callback. A 202 response is acceptance, not completion. Honor Retry-After; "
        "handle failed, cancelled, deduped, unknown and timed-out runs explicitly. Reconcile ambiguous "
        "submissions before retrying. Once ingestion succeeds, an approved external dbt runtime builds "
        "and tests models against the selected Fabric target. Warehouse uses dbt-fabric and writable "
        "SQL; Lakehouse uses dbt-fabricspark through Spark/Livy session jobs, not its read-only SQL "
        "analytics endpoint. The diagram represents a chosen target, not arbitrary automatic movement "
        "between engines. Confirm identity support for the actual item and downstream resources. "
        "Store credentials securely and verify network reachability. Production hosting must be "
        "reliable and always available. No live deployment is implied.\n\nSources:\n"
        + "\n".join([RUN, STATUS, WAREHOUSE, SPARK])
    )


def slide_fabric(prs):
    slide = base(
        prs, 2, "Fabric pipelines orchestrate external dbt",
        "Fabric owns the schedule; an authenticated job service launches the external runtime.",
    )
    panel(slide, 0.5, 7.8, "MICROSOFT FABRIC", "cloud")
    panel(slide, 10.1, 5.4, "CUSTOMER-MANAGED HOST", "runner")
    route(slide, [(3.8, 3.29), (4.6, 3.29)])
    route(slide, [(2.3, 3.85), (2.3, 5.3)])
    route(slide, [(6.15, 3.85), (6.15, 4.3)])
    route(slide, [(6.15, 5.35), (6.15, 5.8)])
    route(slide, [(7.7, 3.1), (10.6, 3.1)])
    route(slide, [(12.8, 3.85), (12.8, 5.3)])
    route(slide, [(7.7, 4.82), (9.2, 4.82), (9.2, 3.62), (10.6, 3.62)],
          color=BLUE, dashed=True)
    route(slide, [(12.8, 6.35), (12.8, 6.72), (2.3, 6.72), (2.3, 6.35)])
    card(slide, 0.8, 2.8, 3.0, 1.05, "Copy activity", "Load Bronze", "copy", size=17)
    card(slide, 4.6, 2.8, 3.1, 1.05, "Web activity", "POST: start dbt", "api", size=17)
    card(slide, 4.6, 4.3, 3.1, 1.05, "Until + Web GET", "Poll external job ID", "clock", size=15.5)
    card(slide, 4.6, 5.8, 3.1, 0.76, "4  Success gate", "Continue or fail", "gate", size=15)
    card(slide, 0.8, 5.3, 3.0, 1.05, "Fabric target", "Warehouse OR Spark", "database", size=16)
    card(slide, 10.6, 2.8, 4.4, 1.05, "Job service API", "Start / status / authentication", "api", size=18)
    card(slide, 10.6, 5.3, 4.4, 1.05, "External dbt runner", "dbt Core + appropriate adapter", "runner", size=18)
    label(slide, "1  Submit", 8.02, 2.7, 2.24)
    label(slide, "2  GET status", 7.87, 5.0, 2.2, color=BLUE)
    label(slide, "Durable job dispatch", 10.75, 4.4, 4.1)
    label(slide, "Bronze", 1.25, 4.45, 2.1)
    label(slide, "3  SQL or Spark/Livy execution", 7.85, 6.45, 4.3, size=11)
    text(slide, "job ID", 6.34, 3.96, 1.0, 0.22, size=10.5, color=MUTED)
    text(slide, "On-premises server or a cloud host", 10.75, 2.61, 4.1, 0.16, size=9, color=MUTED, align=PP_ALIGN.CENTER)
    footer(
        slide,
        "A successful API submission is not a successful transformation.",
        "dbt Core needs a real job service / hosting API. Use an approved network path or gateway where required; never expose a laptop shell.",
        [("Fabric Web activity", WEB, 1.9), ("Warehouse adapter", WAREHOUSE, 1.8),
         ("Lakehouse adapter", SPARK, 1.8)],
    )
    slide.notes_slide.notes_text_frame.text = (
        "Talk track: This reverses the control direction. The Fabric pipeline first loads Bronze "
        "using a Copy activity. On success a Web activity calls the actual external hosting service's "
        "authenticated start API. The service queues or dispatches a durable dbt run and returns an "
        "opaque job ID. Save it, then use bounded Until/Web GET polling or an explicitly implemented "
        "callback protocol to determine the final outcome. The runner reports status through its "
        "service; poll responses return on the same request path. The service, not dbt Core itself, "
        "provides the start/status API. The external runner connects back to the selected Fabric "
        "compute and executes dbt transformations and tests. Warehouse uses dbt-fabric over SQL; "
        "Lakehouse uses dbt-fabricspark over Spark/Livy. Continue downstream only when the external "
        "job explicitly succeeds; fail or escalate failures, cancellation, unknown states and timeout. "
        "Decide retries, duplicate protection, durable correlation, cancellation, credentials and "
        "network reachability before production. If the service is privately reachable, verify the "
        "supported Web activity gateway/network configuration. A gateway is not itself a dbt job "
        "runner. No universal /jobs endpoint, production implementation, or live deployment is implied.\n\nSources:\n"
        + "\n".join([WEB, WAREHOUSE, SPARK,
                     "https://learn.microsoft.com/en-us/fabric/data-factory/webhook-activity"])
    )


def main():
    ICONS.mkdir(parents=True, exist_ok=True)
    for name in ("cloud", "database", "airflow", "runner", "api", "copy", "clock", "gate"):
        icon(name)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    prs.core_properties.title = "Fabric and External Orchestration - Two Architecture Patterns"
    prs.core_properties.subject = "Generic Airflow, dbt and Microsoft Fabric architecture"
    prs.core_properties.author = "Architecture Reference"
    prs.core_properties.keywords = "Fabric, Airflow, dbt, orchestration, architecture"
    slide_external(prs)
    slide_fabric(prs)
    path = ASSETS / "Fabric-External-Orchestration-Architecture.pptx"
    prs.save(path)
    print(f"Created {len(prs.slides)} slides: {path}")


if __name__ == "__main__":
    main()
